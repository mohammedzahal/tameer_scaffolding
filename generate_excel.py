import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from num2words import num2words
import io
import json

STANDARD_FONT = "Arial"

def build_quote_excel(quote_data):
    """
    Builds an Excel workbook (.xlsx) with standardized Arial typography and Tameer Scaffolding branding.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scaffolding Quotation"
    ws.views.sheetView[0].showGridLines = True

    # Standardized Arial typography definitions
    font_company = Font(name=STANDARD_FONT, size=13, bold=True, color="FFFFFF")
    font_subtitle = Font(name=STANDARD_FONT, size=9.5, bold=True, color="FFFFFF")
    font_header = Font(name=STANDARD_FONT, size=9.5, bold=True, color="FFFFFF")
    font_bold = Font(name=STANDARD_FONT, size=9.5, bold=True)
    font_regular = Font(name=STANDARD_FONT, size=9.0)
    font_italic = Font(name=STANDARD_FONT, size=8.5, italic=True)

    fill_brand = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_brand_dark = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    fill_header = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
    fill_totals = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_grand = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    thin_border_side = Side(style='thin', color='CCCCCC')
    dark_border_side = Side(style='thin', color='888888')
    double_bottom_side = Side(style='double', color='000000')

    cell_border = Border(left=dark_border_side, right=dark_border_side, top=dark_border_side, bottom=dark_border_side)
    grand_border = Border(left=dark_border_side, right=dark_border_side, top=dark_border_side, bottom=double_bottom_side)

    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')

    # Data extraction
    company_name = quote_data.get('company_name') or ''
    location = quote_data.get('location') or ''
    attn = quote_data.get('attn') or ''
    yr_ref = quote_data.get('yr_ref') or ''
    quotation_ref = quote_data.get('quotation_ref') or 'TMR-FZ-1748-55471-2026'
    quotation_date = quote_data.get('quotation_date') or ''
    subject = quote_data.get('subject') or 'Scaffolding Materials on sale basis(Used and Refurbed materials)'
    intro_note = quote_data.get('intro_note') or 'We thank you for whatsapp inquiry and pleased quote for used equipment as follows:'
    items = quote_data.get('items') or []
    subtotal = float(quote_data.get('subtotal') or 0.0)
    vat_amount = float(quote_data.get('vat_amount') or 0.0)
    grand_total = float(quote_data.get('grand_total') or 0.0)
    terms = quote_data.get('terms_conditions')
    bank_details = quote_data.get('bank_details')
    sig_name = quote_data.get('signatory_name') or 'Mohamed Faizal'
    sig_title = quote_data.get('signatory_title') or 'Rawaiya AL Etihad Est.'

    # Row 1: Top Brand Banner
    ws.merge_cells('B1:H1')
    cell_title = ws['B1']
    cell_title.value = "مؤسسة تعمير المساحة للمقاولات  |  TAMEER AL MESAHA CONTRACTING EST."
    cell_title.font = font_company
    cell_title.fill = fill_brand
    cell_title.alignment = center_align
    ws.row_dimensions[1].height = 28

    # Row 2: Subtitle & CR
    ws.merge_cells('B2:H2')
    cell_sub = ws['B2']
    cell_sub.value = "C.R.: 2050097965  •  Scaffolding Materials Trading & Services  •  Dammam, KSA"
    cell_sub.font = font_subtitle
    cell_sub.fill = fill_brand_dark
    cell_sub.alignment = center_align
    ws.row_dimensions[2].height = 18

    # Row 4: Client Info Box
    ws['B4'] = "Quotation To:"
    ws['B4'].font = font_bold
    ws['B5'] = f"M/S. {company_name}"
    ws['B5'].font = font_bold
    ws['B6'] = location
    ws['B6'].font = font_regular
    ws['B7'] = f"Attn: {attn}"
    ws['B7'].font = font_regular
    if yr_ref:
        ws['B8'] = f"Yr. Ref: {yr_ref}"
        ws['B8'].font = font_italic

    # Right side: Ref & Date
    ws['F4'] = "Quotation Ref:"
    ws['F4'].font = font_bold
    ws['G4'] = quotation_ref
    ws['G4'].font = font_bold

    ws['F5'] = "Date:"
    ws['F5'].font = font_bold
    ws['G5'] = quotation_date
    ws['G5'].font = font_regular

    ws['F6'] = "Currency:"
    ws['F6'].font = font_bold
    ws['G6'] = "Saudi Riyals (SAR)"
    ws['G6'].font = font_regular

    # Row 10: Subject & Intro
    ws['B10'] = f"Subject: {subject}"
    ws['B10'].font = font_bold

    ws['B11'] = intro_note
    ws['B11'].font = font_italic

    # Row 13: Table Headers
    headers = ["SL", "Item Description", "Unit", "Qty", "Unit Rate (SR)", "Amount (SR)", "Remarks"]
    cols = ['B', 'C', 'D', 'E', 'F', 'G', 'H']
    header_row = 13
    ws.row_dimensions[header_row].height = 22

    for col_letter, h in zip(cols, headers):
        c = ws[f"{col_letter}{header_row}"]
        c.value = h
        c.font = font_header
        c.fill = fill_header
        c.alignment = center_align
        c.border = cell_border

    # Item Rows
    current_row = 14
    for idx, item in enumerate(items, 1):
        ws.row_dimensions[current_row].height = 19
        qty = float(item.get('qty', 0))
        rate = float(item.get('rate', 0.0))
        amt = float(item.get('amount', qty * rate))

        ws[f"B{current_row}"] = idx
        ws[f"B{current_row}"].alignment = center_align

        ws[f"C{current_row}"] = item.get('desc', '')
        ws[f"C{current_row}"].alignment = left_align

        ws[f"D{current_row}"] = item.get('unit', 'Pcs.')
        ws[f"D{current_row}"].alignment = center_align

        ws[f"E{current_row}"] = qty
        ws[f"E{current_row}"].alignment = right_align
        ws[f"E{current_row}"].number_format = '#,##0'

        ws[f"F{current_row}"] = rate
        ws[f"F{current_row}"].alignment = right_align
        ws[f"F{current_row}"].number_format = '#,##0.00'

        ws[f"G{current_row}"] = amt
        ws[f"G{current_row}"].alignment = right_align
        ws[f"G{current_row}"].number_format = '#,##0.00'

        ws[f"H{current_row}"] = item.get('remarks', '')
        ws[f"H{current_row}"].alignment = left_align

        for col_letter in cols:
            cell = ws[f"{col_letter}{current_row}"]
            cell.font = font_regular
            cell.border = cell_border

        current_row += 1

    # Totals Section
    ws.row_dimensions[current_row].height = 20
    ws.merge_cells(f'B{current_row}:F{current_row}')
    ws[f'B{current_row}'] = "Total SR (Subtotal)"
    ws[f'B{current_row}'].font = font_bold
    ws[f'B{current_row}'].alignment = right_align
    ws[f'G{current_row}'] = subtotal
    ws[f'G{current_row}'].font = font_bold
    ws[f'G{current_row}'].alignment = right_align
    ws[f'G{current_row}'].number_format = '#,##0.00'
    for col_letter in cols:
        ws[f'{col_letter}{current_row}'].fill = fill_totals
        ws[f'{col_letter}{current_row}'].border = cell_border
    current_row += 1

    # VAT 15%
    ws.row_dimensions[current_row].height = 20
    ws.merge_cells(f'B{current_row}:F{current_row}')
    ws[f'B{current_row}'] = "VAT @ 15%"
    ws[f'B{current_row}'].font = font_bold
    ws[f'B{current_row}'].alignment = right_align
    ws[f'G{current_row}'] = vat_amount
    ws[f'G{current_row}'].font = font_bold
    ws[f'G{current_row}'].alignment = right_align
    ws[f'G{current_row}'].number_format = '#,##0.00'
    for col_letter in cols:
        ws[f'{col_letter}{current_row}'].fill = fill_totals
        ws[f'{col_letter}{current_row}'].border = cell_border
    current_row += 1

    # Grand Total
    ws.row_dimensions[current_row].height = 22
    ws.merge_cells(f'B{current_row}:F{current_row}')
    ws[f'B{current_row}'] = "Grand Total (SR)"
    ws[f'B{current_row}'].font = font_bold
    ws[f'B{current_row}'].alignment = right_align
    ws[f'G{current_row}'] = grand_total
    ws[f'G{current_row}'].font = font_bold
    ws[f'G{current_row}'].alignment = right_align
    ws[f'G{current_row}'].number_format = '#,##0.00'
    for col_letter in cols:
        ws[f'{col_letter}{current_row}'].fill = fill_grand
        ws[f'{col_letter}{current_row}'].border = grand_border
    current_row += 1

    # Total in words
    current_row += 1
    try:
        words = num2words(grand_total, to='currency', lang='en', currency='SAR')
        ws.merge_cells(f'B{current_row}:H{current_row}')
        ws[f'B{current_row}'] = f"Amount in Words: {words.title()}"
        ws[f'B{current_row}'].font = font_bold
    except:
        pass
    current_row += 2

    # Terms & Conditions Box
    ws['B{0}'.format(current_row)] = "Terms & Conditions:"
    ws['B{0}'.format(current_row)].font = font_bold
    current_row += 1

    if terms:
        if isinstance(terms, str):
            terms_lines = [l.strip() for l in terms.split('\n') if l.strip()]
        elif isinstance(terms, list):
            terms_lines = terms
        else:
            terms_lines = []
        for line in terms_lines:
            ws['B{0}'.format(current_row)] = line
            ws['B{0}'.format(current_row)].font = font_regular
            current_row += 1
    current_row += 1

    # Bank Details Box
    ws['B{0}'.format(current_row)] = "Banking Details:"
    ws['B{0}'.format(current_row)].font = font_bold
    current_row += 1

    if bank_details:
        if isinstance(bank_details, str):
            bank_lines = [l.strip() for l in bank_details.split('\n') if l.strip()]
        elif isinstance(bank_details, list):
            bank_lines = bank_details
        else:
            bank_lines = []
        for line in bank_lines:
            ws['B{0}'.format(current_row)] = line
            ws['B{0}'.format(current_row)].font = font_regular
            current_row += 1
    current_row += 1

    # Closing & Signatory
    ws['B{0}'.format(current_row)] = "Thank You and Best regards,"
    ws['B{0}'.format(current_row)].font = font_regular
    current_row += 2

    ws['B{0}'.format(current_row)] = sig_name
    ws['B{0}'.format(current_row)].font = font_bold
    current_row += 1

    ws['B{0}'.format(current_row)] = sig_title
    ws['B{0}'.format(current_row)].font = font_regular

    col_widths = {
        'A': 4,
        'B': 7,
        'C': 38,
        'D': 10,
        'E': 12,
        'F': 16,
        'G': 18,
        'H': 18
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
