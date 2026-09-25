import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
import copy
import io
import os
import re

TEMPLATE_PATH = r'D:\TAMEER\Quotations\Word\quotation ref.docx'

STANDARD_FONT = "Arial"

def set_run_font(run, font_name=STANDARD_FONT, size_pt=9.5, bold=False, color="000000"):
    """Standardizes typography for a Word run across all versions."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)

def format_num(val):
    try:
        return f"{float(val):,.2f}"
    except:
        return str(val)

def build_quote_word(quote_data):
    """
    Builds a Word document (.docx) using standardized Arial typography and official Tameer layout.
    """
    doc = docx.Document(TEMPLATE_PATH)

    company_name = quote_data.get('company_name', '')
    location = quote_data.get('location', 'Kingdom of Saudi Arabia')
    attn = quote_data.get('attn', '')
    yr_ref = quote_data.get('yr_ref', '')
    quotation_ref = quote_data.get('quotation_ref', 'TMR-FZ-1748-55471-2026')
    quotation_date = quote_data.get('quotation_date', '24 Sep 2026')
    subject = quote_data.get('subject', 'Scaffolding Materials on sale basis(Used and Refurbed materials)')
    intro_note = quote_data.get('intro_note', 'We thank you for whatsapp inquiry and pleased quote for used equipment as follows:')
    items = quote_data.get('items', [])
    subtotal = float(quote_data.get('subtotal', 0.0))
    vat_amount = float(quote_data.get('vat_amount', 0.0))
    grand_total = float(quote_data.get('grand_total', 0.0))

    # 1. Update Date and Quotation Ref in Top-Right Textbox on Page 1
    # Standardized to Date and Ref only (removes Page 1 of 2 indicator)
    p1 = doc.paragraphs[1]
    for p_inner in p1._p.findall('.//{*}p'):
        p_text = ''.join(p_inner.itertext())
        if 'Page' in p_text:
            for t in p_inner.findall('.//{*}t'):
                t.text = ''
        else:
            for r in p_inner.findall('.//{*}r'):
                for t in r.findall('.//{*}t'):
                    if t.text:
                        if 'July 09, 2023' in t.text or 'Date:' in t.text:
                            t.text = f"Date: {quotation_date}"
                        elif 'REE-07-07/23' in t.text or 'Ref' in t.text:
                            t.text = f"Our Ref: {quotation_ref}"
                # Apply standardized font to inner textbox runs
                r_elem = docx.text.run.Run(r, p_inner)
                set_run_font(r_elem, STANDARD_FONT, 9.5, bold=True)

    # 2. Update Client Details paragraphs with standardized Arial typography
    if len(doc.paragraphs) > 2:
        doc.paragraphs[2].text = f"M/S. {company_name}"
        if doc.paragraphs[2].runs:
            set_run_font(doc.paragraphs[2].runs[0], STANDARD_FONT, 10.0, bold=True)

    if len(doc.paragraphs) > 3:
        doc.paragraphs[3].text = location
        if doc.paragraphs[3].runs:
            set_run_font(doc.paragraphs[3].runs[0], STANDARD_FONT, 9.5, bold=True)

    if len(doc.paragraphs) > 4:
        doc.paragraphs[4].text = f"Attn:.: {attn}"
        if doc.paragraphs[4].runs:
            set_run_font(doc.paragraphs[4].runs[0], STANDARD_FONT, 9.5, bold=True)

    if len(doc.paragraphs) > 5:
        doc.paragraphs[5].text = f"Yr. reference : {yr_ref}" if yr_ref else "Yr. reference : "
        if doc.paragraphs[5].runs:
            set_run_font(doc.paragraphs[5].runs[0], STANDARD_FONT, 9.5, bold=True)

    if len(doc.paragraphs) > 6:
        doc.paragraphs[6].text = f"Sub:   {subject}"
        if doc.paragraphs[6].runs:
            set_run_font(doc.paragraphs[6].runs[0], STANDARD_FONT, 10.0, bold=True)

    if len(doc.paragraphs) > 8:
        doc.paragraphs[8].text = intro_note
        if doc.paragraphs[8].runs:
            set_run_font(doc.paragraphs[8].runs[0], STANDARD_FONT, 9.5, bold=False)

    # 3. Standardize Items Table
    table = doc.tables[0]
    num_items = len(items)
    template_item_count = 24

    def format_cell(cell, text, size=9.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = str(text)
        if cell.paragraphs:
            p = cell.paragraphs[0]
            p.alignment = align
            if p.runs:
                set_run_font(p.runs[0], STANDARD_FONT, size, bold=bold)

    # Standardize header row (row 0)
    header_alignments = [
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER
    ]
    for col_idx, cell in enumerate(table.rows[0].cells):
        if cell.paragraphs and cell.paragraphs[0].runs:
            cell.paragraphs[0].alignment = header_alignments[col_idx]
            set_run_font(cell.paragraphs[0].runs[0], STANDARD_FONT, 9.0, bold=True)

    # Adjust row count
    if num_items < template_item_count:
        for _ in range(template_item_count - num_items):
            tr = table.rows[num_items + 1]._tr
            tr.getparent().remove(tr)
    elif num_items > template_item_count:
        sample_tr = copy.deepcopy(table.rows[1]._tr)
        totals_tr = table.rows[-3]._tr
        for _ in range(num_items - template_item_count):
            new_tr = copy.deepcopy(sample_tr)
            totals_tr.addprevious(new_tr)

    # Populate items with clean standardized typography
    for idx, item in enumerate(items):
        r = table.rows[idx + 1]
        format_cell(r.cells[0], str(item.get('sl', idx + 1)), size=9.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(r.cells[1], item.get('desc', ''), size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(r.cells[2], item.get('unit', 'Pcs.'), size=9.0, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(r.cells[3], f"{item.get('qty', 0):,}", size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        format_cell(r.cells[4], format_num(item.get('rate', 0.0)), size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        format_cell(r.cells[5], format_num(item.get('amount', 0.0)), size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        format_cell(r.cells[6], item.get('remarks', ''), size=8.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)

    # Populate Totals with standardized font
    format_cell(table.rows[-3].cells[5], format_num(subtotal), size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    format_cell(table.rows[-2].cells[5], format_num(vat_amount), size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    format_cell(table.rows[-1].cells[5], format_num(grand_total), size=10.0, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # 4. Remove all intermediate empty paragraphs between Table and Terms & Conditions
    terms_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if 'Terms & Conditions' in p.text:
            terms_idx = i
            break

    if terms_idx > 9:
        paras_to_delete = [doc.paragraphs[k] for k in range(9, terms_idx)]
        for p in paras_to_delete:
            p._p.getparent().remove(p._p)

    # Insert a clean Word Page Break right before Terms & Conditions so it cleanly starts on Page 2
    terms_para = None
    for p in doc.paragraphs:
        if 'Terms & Conditions' in p.text:
            terms_para = p
            break

    if terms_para:
        break_para = terms_para.insert_paragraph_before()
        break_para.add_run().add_break(WD_BREAK.PAGE)
        # Style the heading
        if terms_para.runs:
            set_run_font(terms_para.runs[0], STANDARD_FONT, 10.5, bold=True)
            terms_para.runs[0].underline = True

    # 5. Standardize Terms & Conditions Typography
    terms = quote_data.get('terms_conditions', '')
    if isinstance(terms, list):
        terms_lines = terms
    else:
        terms_lines = [line.strip() for line in terms.split('\n') if line.strip()]

    current_terms_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if 'Terms & Conditions' in p.text:
            current_terms_idx = i
            break

    if current_terms_idx != -1:
        for offset, t_line in enumerate(terms_lines[:4]):
            target_idx = current_terms_idx + 1 + offset
            if target_idx < len(doc.paragraphs):
                p_target = doc.paragraphs[target_idx]
                clean_text = re.sub(r'^\s*(\d+[\.\)]\s*)+', '', t_line)
                p_target.text = clean_text
                if p_target.runs:
                    set_run_font(p_target.runs[0], STANDARD_FONT, 9.5, bold=(offset == 0))

    # 6. Standardize Bank Details Typography
    bank_text = quote_data.get('bank_details', '')
    if bank_text:
        bank_lines = [l.strip() for l in bank_text.split('\n') if l.strip()]
        for p in doc.paragraphs:
            if 'Saudi National Bank' in p.text and len(bank_lines) > 0:
                first_bank_line = re.sub(r'^(Our\s*:\s*)+', '', bank_lines[0])
                first_bank_line = re.sub(r'^(Bank\s*:\s*)+', '', first_bank_line)
                p.text = f"Our : Bank : {first_bank_line}"
                if p.runs: set_run_font(p.runs[0], STANDARD_FONT, 9.5, bold=True)
            elif 'Account #' in p.text and len(bank_lines) > 1:
                p.text = f"\t{bank_lines[1]}"
                if p.runs: set_run_font(p.runs[0], STANDARD_FONT, 9.5, bold=True)
            elif 'IBAN' in p.text and len(bank_lines) > 2:
                p.text = bank_lines[2]
                if p.runs: set_run_font(p.runs[0], STANDARD_FONT, 9.5, bold=True)

    # 7. Standardize Courtesy & Signatory Typography
    for p in doc.paragraphs:
        if 'Should you require' in p.text:
            if p.runs: set_run_font(p.runs[0], STANDARD_FONT, 9.0, bold=False)
        elif 'Thank You and Best regards' in p.text:
            if p.runs: set_run_font(p.runs[0], STANDARD_FONT, 9.5, bold=False)
        elif 'Mohamed Faizal' in p.text:
            sig_name = quote_data.get('signatory_name', 'Mohamed Faizal')
            p.text = sig_name
            if p.runs: set_run_font(p.runs[0], STANDARD_FONT, 10.0, bold=True)
        elif 'Rawaiya AL Etihad Est.' in p.text:
            sig_title = quote_data.get('signatory_title', 'Rawaiya AL Etihad Est.')
            p.text = sig_title
            if p.runs: set_run_font(p.runs[0], STANDARD_FONT, 9.5, bold=True)

    output_stream = io.BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)
    return output_stream
